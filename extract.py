# -*- coding: utf-8 -*-
"""Extract all study materials (PDF/DOCX/XLS) to plain text for study."""
import os, sys, traceback
import fitz  # pymupdf
import docx as pydocx

ROOT = r"D:\project\writingzcode"
OUT = os.path.join(ROOT, "extracted")
os.makedirs(OUT, exist_ok=True)

def safe_name(rel):
    return rel.replace("\\", "__").replace("/", "__") + ".txt"

def extract_pdf(src, dst):
    doc = fitz.open(src)
    parts = []
    for i, page in enumerate(doc):
        t = page.get_text("text").strip()
        if t:
            parts.append(f"\n===== [page {i+1}] =====\n{t}")
    doc.close()
    with open(dst, "w", encoding="utf-8") as f:
        f.write("\n".join(parts))
    return sum(len(p) for p in parts)

def extract_docx(src, dst):
    d = pydocx.Document(src)
    lines = []
    for p in d.paragraphs:
        if p.text.strip():
            lines.append(p.text)
    for tbl in d.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    text = "\n".join(lines)
    with open(dst, "w", encoding="utf-8") as f:
        f.write(text)
    return len(text)

def extract_xls(src, dst):
    try:
        import xlrd
        wb = xlrd.open_workbook(src)
        lines = []
        for sh in wb.sheets():
            lines.append(f"\n===== sheet: {sh.name} =====")
            for r in range(sh.nrows):
                vals = [str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)]
                if any(vals):
                    lines.append(" | ".join(vals))
        text = "\n".join(lines)
        with open(dst, "w", encoding="utf-8") as f:
            f.write(text)
        return len(text)
    except Exception as e:
        return -1

results = []
for dirpath, dirnames, filenames in os.walk(ROOT):
    dirnames[:] = [d for d in dirnames if d not in ("extracted", "node_modules", ".git", "tool", "knowledge")]
    for fn in filenames:
        src = os.path.join(dirpath, fn)
        low = fn.lower()
        rel = os.path.relpath(src, ROOT)
        if low.endswith(".pdf"):
            dst = os.path.join(OUT, safe_name(rel))
            try:
                n = extract_pdf(src, dst)
                results.append((rel, "pdf", n))
            except Exception as e:
                results.append((rel, "pdf", f"FAIL {e}"))
        elif low.endswith(".docx"):
            dst = os.path.join(OUT, safe_name(rel))
            try:
                n = extract_docx(src, dst)
                results.append((rel, "docx", n))
            except Exception as e:
                results.append((rel, "docx", f"FAIL {e}"))
        elif low.endswith(".xls"):
            dst = os.path.join(OUT, safe_name(rel))
            n = extract_xls(src, dst)
            results.append((rel, "xls", n))

with open(os.path.join(OUT, "_index.txt"), "w", encoding="utf-8") as f:
    for rel, kind, n in sorted(results):
        f.write(f"{kind}\t{n}\t{rel}\n")
print("done", len(results))
for rel, kind, n in sorted(results):
    print(f"{kind}\t{n}\t{rel}")
